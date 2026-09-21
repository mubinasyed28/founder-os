import io
import json
import logging
from typing import Dict, Any, Tuple
from backend.config import settings

logger = logging.getLogger(__name__)

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    ext = filename.split(".")[-1].lower()
    
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join([page.extract_text() or "" for page in reader.pages])
        return text
        
    elif ext in ["docx", "doc"]:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join([p.text for p in doc.paragraphs])
        return text
        
    elif ext in ["pptx", "ppt"]:
        import pptx
        prs = pptx.Presentation(io.BytesIO(file_bytes))
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)
        
    elif ext in ["txt", "md", "markdown"]:
        return file_bytes.decode("utf-8", errors="ignore")
        
    else:
        raise ValueError(f"Unsupported file format: .{ext}. Supported formats: PDF, DOCX, PPTX, MD, TXT.")

async def process_document_and_extract(
    file_bytes: bytes,
    filename: str,
    current_profile: Dict[str, Any]
) -> Tuple[str, Dict[str, Any]]:
    """Extracts text from uploaded document and asks Mistral AI to extract startup details."""
    extracted_text = extract_text_from_file(file_bytes, filename)
    
    apiKey = settings.MISTRAL_API_KEY
    if apiKey:
        try:
            import asyncio
            import urllib.request
            
            prompt = (
                f"You are an AI Co-Founder analyzing an uploaded document for a startup.\n"
                f"Document Name: {filename}\n"
                f"Document Content Snippet (first 4000 chars):\n{extracted_text[:4000]}\n\n"
                f"Current Known Profile State:\n{json.dumps(current_profile, indent=2)}\n\n"
                f"Extract any new or updated startup details from this document and return a raw JSON object."
            )

            schema_desc = {
                "acknowledgement": "A natural 2-sentence confirmation explaining what you learned from the document...",
                "profile_updates": {
                    "company_name": "string", "one_line_pitch": "string", "description": "string", 
                    "problem_statement": "string", "solution": "string", "product": "string", 
                    "target_customer": "string", "industry": "string", "business_model": "string", 
                    "revenue_model": "string", "current_stage": "string", "founders": ["string"], 
                    "team_size": "string", "website": "string", "location": "string", 
                    "current_progress": "string", "validation_done": "boolean", "validation_details": "string", 
                    "customer_interviews": "string", "survey_details": "string", "traction": "string", 
                    "competitors": ["string"], "current_challenges": ["string"], "short_term_goals": ["string"], 
                    "long_term_vision": "string"
                }
            }

            def call_mistral():
                candidate_models = ["mistral-small-latest", "ministral-8b-latest", "open-mistral-7b", "ministral-3b-latest"]
                last_err = None
                for model_name in candidate_models:
                    payload = {
                        "model": model_name,
                        "messages": [
                            {
                                "role": "system",
                                "content": f"You are a structured data extractor. Return JSON matching this format: {json.dumps(schema_desc)}"
                            },
                            {
                                "role": "user",
                                "content": prompt
                            }
                        ],
                        "response_format": {
                            "type": "json_object"
                        },
                        "max_tokens": 1500
                    }
                    
                    headers = {
                        "Content-Type": "application/json",
                        "Authorization": f"Bearer {apiKey}"
                    }
                    
                    req = urllib.request.Request(
                        "https://api.mistral.ai/v1/chat/completions",
                        data=json.dumps(payload).encode("utf-8"),
                        headers=headers,
                        method="POST"
                    )
                    
                    try:
                        with urllib.request.urlopen(req, timeout=40) as response:
                            resp_data = response.read().decode("utf-8")
                            resp_json = json.loads(resp_data)
                            content = resp_json["choices"][0]["message"]["content"]
                            return json.loads(content.strip())
                    except Exception as e:
                        logger.warning(f"[doc_service] Model {model_name} failed: {e}")
                        last_err = e
                        continue
                if last_err:
                    raise last_err
                raise RuntimeError("All Mistral doc extraction candidate models failed")

            parsed = await asyncio.to_thread(call_mistral)
            ack = parsed.get("acknowledgement", f"I've read through {filename} and extracted key details into your company profile!")
            updates = parsed.get("profile_updates", {})
            return ack, updates
        except Exception as e:
            logger.warning(f"[doc_service] Mistral document extraction failed: {e}")

    ack = f"I've parsed {filename} ({len(extracted_text)} characters) and added key information to your profile."
    updates = {"description": extracted_text[:300]}
    return ack, updates

